# -*- coding: utf-8 -*-
"""Round-5: sync Figure-10 numbers with the embedded figure legend (0.489), soften the band claim."""
import copy, os
from docx import Document
from docx.shared import RGBColor

BASE = os.path.dirname(os.path.abspath(__file__))
DOC = os.path.join(os.path.dirname(BASE), "YOLO-CAST_修改稿(2).docx")
BLUE = RGBColor(0x00, 0x00, 0xFF)
d = Document(DOC)

def base_rpr(p):
    return copy.deepcopy(p.runs[0]._element.get_or_add_rPr()) if p.runs else None

# ---- P97 ----
p97 = d.paragraphs[97]
full = p97.text
old = ("peaks at 0.94 at a confidence threshold of 0.538 and stays above 0.85 for thresholds between "
       "0.02 and 0.90, which covers nearly the entire range")
new = "peaks at 0.94 at a confidence threshold of 0.489 and stays above 0.85 across a wide range of thresholds"
assert old in full, full[:120]
full = full.replace(old, new)
frags = [
    "97.6% mAP@0.5",
    "Leaf Rust (88.5%) and Sharp Eyespot (96.1%)",
    new,
]
rpr = base_rpr(p97)
for r in list(p97.runs):
    r._element.getparent().remove(r._element)
rest = full
for fr in frags:
    idx = rest.index(fr)
    for text, blue in ((rest[:idx], False), (fr, True)):
        if text:
            run = p97.add_run(text)
            if rpr is not None:
                run._element.insert(0, copy.deepcopy(rpr))
            if blue:
                run.font.color.rgb = BLUE
    rest = rest[idx + len(fr):]
if rest:
    run = p97.add_run(rest)
    if rpr is not None:
        run._element.insert(0, copy.deepcopy(rpr))

# ---- P101 caption ----
p101 = d.paragraphs[101]
full = p101.text
old = "(all classes: peak F1 = 0.94 at a confidence of 0.538)"
new = "(all classes: peak F1 = 0.94 at a confidence of 0.489)"
assert old in full, full[:100]
full = full.replace(old, new)
pre, post = full.split(new, 1)
rpr = base_rpr(p101)
for r in list(p101.runs):
    r._element.getparent().remove(r._element)
for text, blue in ((pre, False), (new, True), (post, False)):
    if text:
        run = p101.add_run(text)
        if rpr is not None:
            run._element.insert(0, copy.deepcopy(rpr))
        if blue:
            run.font.color.rgb = BLUE

d.save(DOC)
print("EDIT OK")
d2 = Document(DOC)
print("P97:", d2.paragraphs[97].text[:330])
print()
print("P101:", d2.paragraphs[101].text)
